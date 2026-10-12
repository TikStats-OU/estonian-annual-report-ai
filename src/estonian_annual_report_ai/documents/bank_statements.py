"""Deterministic importers for common bank statement formats."""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from itertools import pairwise
from pathlib import Path

from .common import (
    MAX_EXTRACTED_TEXT_CHARS,
    MAX_PDF_PAGES,
    DocumentParseError,
    DocumentValidationError,
    decode_text,
    normalize_currency,
    normalize_label,
    parse_money,
    validate_document,
)
from .models import BankTransaction, ImportIssue, ParsedBankStatement, SourceReference

_HEADERS: dict[str, set[str]] = {
    "transaction_date": {
        "date", "transaction date", "booking date", "value date", "payment date",
        "kuupaev", "tehingu kuupaev", "vaartuspaev", "makse kuupaev",
    },
    "amount": {"amount", "transaction amount", "sum", "summa", "tehingu summa", "makse summa"},
    "debit": {"debit", "debit amount", "withdrawal", "deebet", "deebet summa", "valja minek"},
    "credit": {"credit", "credit amount", "deposit", "kreedit", "kreedit summa", "sisse tulek"},
    "currency": {"currency", "currency code", "valuuta", "valuuta kood"},
    "counterparty": {
        "counterparty", "beneficiary", "payer", "payee", "sender", "recipient",
        "payer name", "payee name",
        "saaja", "maksja", "tehingu osapool", "partner", "name",
    },
    "description": {
        "description", "details", "payment details", "remittance information", "purpose",
        "selgitus", "makse selgitus", "tehingu kirjeldus", "kirjeldus", "details of payment",
    },
    "reference": {
        "reference", "transaction id", "transaction number", "reference number", "payment reference",
        "transaction reference", "transferwise id", "viitenumber", "tehingu number", "tehingu id", "document number",
    },
    "direction": {"direction", "type", "flow", "transaction type", "suund", "tehingu liik"},
    "opening_balance": {"opening balance", "starting balance", "balance brought forward", "algsaldo", "algbilanss"},
    "closing_balance": {"closing balance", "ending balance", "balance carried forward", "lopp saldo", "loppbilanss", "lopp saldo kokku"},
    "account_id": {"iban", "account", "account number", "bank account", "kontonumber", "konto number"},
    "balance": {"balance", "running balance", "saldo", "account balance"},
    "merchant": {"merchant"},
}

_DATE_FORMATS = (
    "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y",
    "%d.%m.%y", "%d/%m/%y", "%d-%m-%y",
)
_DATE_AT_START = re.compile(r"^\s*(\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[./-]\d{1,2}[./-]\d{2,4})\b")
_BALANCE_LABELS = {
    "opening": ("opening balance", "starting balance", "balance brought forward", "algsaldo", "algbilanss"),
    "closing": ("closing balance", "ending balance", "balance carried forward", "lopp saldo", "loppbilanss"),
}
_DIRECTION_ALIASES = {
    "debit": "debit", "dr": "debit", "withdrawal": "debit", "outgoing": "debit", "deebet": "debit",
    "credit": "credit", "cr": "credit", "deposit": "credit", "incoming": "credit", "kreedit": "credit",
}
_REFERENCE_LABEL = re.compile(r"(?i)\b(?:reference|ref\.?|viitenumber|payment ref(?:erence)?)\s*[:#]?\s*([\w/-]+)")
_ACCOUNT_LABEL = re.compile(r"(?i)\b(?:IBAN|account number|bank account|kontonumber)\s*[:#]?\s*([A-Z]{2}[A-Z0-9 ]{10,32}|[A-Z0-9-]{5,32})")
_CURRENCY_CODE = r"(?:[A-Z]{3}|[€£¥$])"


def parse_bank_statement(path: str | Path, *, max_bytes: int | None = None) -> ParsedBankStatement:
    """Parse a local CSV, XLSX, or text-based PDF bank statement."""
    kwargs = {} if max_bytes is None else {"max_bytes": max_bytes}
    document = validate_document(path, **kwargs)
    extension = Path(path).suffix.lower()
    if extension == ".csv":
        return _parse_csv(document.data, document.source)
    if extension == ".xlsx":
        return _parse_xlsx(document.data, document.source)
    if extension == ".pdf":
        return _parse_pdf(document.data, document.source)
    raise DocumentValidationError("Bank statements must be PDF, CSV, or XLSX files.")


def _parse_csv(data: bytes, source: SourceReference) -> ParsedBankStatement:
    text = decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    try:
        for row_number, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
            if row_number > 250_000:
                raise DocumentValidationError("CSV exceeds the row limit.")
            rows.append(row)
    except csv.Error as exc:
        raise DocumentParseError("CSV document could not be parsed safely.") from exc
    return _parse_rows(((row, source.at(row_number=index)) for index, row in enumerate(rows, 1)), source)


def _parse_xlsx(data: bytes, source: SourceReference) -> ParsedBankStatement:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParseError("XLSX parsing requires the openpyxl package.") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        statements: list[ParsedBankStatement] = []
        total_rows = 0
        for sheet in workbook.worksheets:
            if sheet.max_row and sheet.max_row > 250_000:
                raise DocumentValidationError("XLSX worksheet exceeds the row limit.")
            row_records: list[tuple[Sequence[object], SourceReference]] = []
            for row_number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                if any(value is not None and str(value).strip() for value in cells):
                    total_rows += 1
                    if total_rows > 500_000:
                        raise DocumentValidationError("XLSX workbook exceeds the row limit.")
                    row_records.append((cells, source.at(row_number=row_number, sheet_name=sheet.title)))
            if row_records:
                statements.append(_parse_rows(iter(row_records), source, emit_no_header=False))
        workbook.close()
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentParseError("XLSX workbook could not be parsed safely.") from exc
    transactions: list[BankTransaction] = []
    issues: list[ImportIssue] = []
    balances: dict[str, tuple[Decimal, str | None]] = {}
    account_id: str | None = None
    currencies: set[str] = set()
    for statement in statements:
        transactions.extend(statement.transactions)
        issues.extend(statement.issues)
        account_id = account_id or statement.bank_account_id
        if statement.currency:
            currencies.add(statement.currency)
        _merge_balances(balances, statement, issues, source)
    if not transactions and statements:
        issues.append(ImportIssue("missing_transaction_header", "No recognizable transaction table header was found.", source))
    currency = next(iter(currencies)) if len(currencies) == 1 else None
    return ParsedBankStatement(
        source=source,
        transactions=tuple(transactions),
        opening_balance=balances.get("opening", (None, None))[0],
        closing_balance=balances.get("closing", (None, None))[0],
        currency=currency,
        bank_account_id=account_id,
        issues=tuple(issues),
    )


def _parse_pdf(data: bytes, source: SourceReference) -> ParsedBankStatement:
    try:
        import pdfplumber
    except ImportError as exc:
        raise DocumentParseError("PDF parsing requires the pdfplumber package.") from exc

    transactions: list[BankTransaction] = []
    issues: list[ImportIssue] = []
    balances: dict[str, tuple[Decimal, str | None]] = {}
    account_id: str | None = None
    explicit_currencies: set[str] = set()
    found_tables = False
    has_extractable_content = False
    extracted_text_chars = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise DocumentValidationError("PDF exceeds the page limit.")
            for page_number, page in enumerate(pdf.pages, 1):
                page_ref = source.at(page_number=page_number)
                page_text = page.extract_text() or ""
                has_extractable_content = has_extractable_content or bool(page_text.strip())
                extracted_text_chars += len(page_text)
                if extracted_text_chars > MAX_EXTRACTED_TEXT_CHARS:
                    raise DocumentValidationError("PDF extracted text exceeds the safety limit.")
                lines = page_text.splitlines()
                for line_number, line in enumerate(lines, 1):
                    _capture_text_metadata(line, line_number, page_ref, balances, issues)
                    if account_id is None:
                        match = _ACCOUNT_LABEL.search(line)
                        if match:
                            account_id = re.sub(r"\s+", "", match.group(1)).upper()

                page_had_transactions = False
                page_tables = page.extract_tables() or []
                has_extractable_content = has_extractable_content or bool(page_tables)
                for table in page_tables:
                    records = [(row, page_ref.at(row_number=row_index)) for row_index, row in enumerate(table, 1) if row]
                    parsed = _parse_rows(iter(records), source, emit_no_header=False)
                    if parsed.transactions:
                        found_tables = True
                        page_had_transactions = True
                        transactions.extend(parsed.transactions)
                        issues.extend(parsed.issues)
                        _merge_balances(balances, parsed, issues, page_ref)
                        explicit_currencies.update(
                            transaction.currency for transaction in parsed.transactions if transaction.currency
                        )
                        account_id = account_id or parsed.bank_account_id
                if not page_had_transactions:
                    for line_number, line in enumerate(lines, 1):
                        transaction = _parse_pdf_transaction(line, page_ref.at(row_number=line_number), account_id, issues)
                        if transaction is not None:
                            transactions.append(transaction)
                            if transaction.currency:
                                explicit_currencies.add(transaction.currency)
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentParseError("PDF document could not be parsed safely.") from exc

    if not transactions:
        code = "no_transactions" if has_extractable_content or found_tables else "no_text_content"
        message = "PDF contains no extractable text or tables." if code == "no_text_content" else "No transactions were recognized."
        issues.append(ImportIssue(code, message, source))
    currency = next(iter(explicit_currencies)) if len(explicit_currencies) == 1 else None
    if not currency:
        for _, value_currency in balances.values():
            if value_currency:
                if currency is None:
                    currency = value_currency
                elif currency != value_currency:
                    currency = None
                    break
    transactions = [
        transaction if transaction.currency else BankTransaction(
            transaction_date=transaction.transaction_date,
            amount=transaction.amount,
            currency=currency,
            source=transaction.source,
            bank_account_id=transaction.bank_account_id or account_id,
            counterparty=transaction.counterparty,
            description=transaction.description,
            reference=transaction.reference,
            direction=transaction.direction,
        )
        for transaction in transactions
    ]
    return ParsedBankStatement(
        source=source,
        transactions=tuple(transactions),
        opening_balance=balances.get("opening", (None, None))[0],
        closing_balance=balances.get("closing", (None, None))[0],
        currency=currency,
        bank_account_id=account_id,
        issues=tuple(issues),
    )


def _parse_rows(
    records: Iterable[tuple[Sequence[object], SourceReference]],
    source: SourceReference,
    *,
    emit_no_header: bool = True,
) -> ParsedBankStatement:
    materialized = list(records)
    header_index: int | None = None
    columns: dict[str, int] = {}
    for index, (row, _) in enumerate(materialized[:30]):
        candidate = _find_columns(row)
        if "transaction_date" in candidate and ("amount" in candidate or "debit" in candidate or "credit" in candidate):
            header_index, columns = index, candidate
            break

    issues: list[ImportIssue] = []
    transactions: list[BankTransaction] = []
    balances: dict[str, tuple[Decimal, str | None]] = {}
    account_id: str | None = None
    transaction_account_ids: set[str] = set()
    currencies: set[str] = set()
    running_rows: list[tuple[date, Decimal, Decimal, SourceReference]] = []

    for row, ref in materialized:
        row_columns = _find_columns(row)
        if "transaction_date" in row_columns and (
            "amount" in row_columns or "debit" in row_columns or "credit" in row_columns
        ):
            continue
        for kind, labels in _BALANCE_LABELS.items():
            joined = " ".join(str(value or "") for value in row)
            if any(normalize_label(label) in normalize_label(joined) for label in labels):
                amount, currency = _find_amount_in_row(row)
                if amount is not None:
                    _store_balance(balances, kind, amount, currency, ref, issues)
        if account_id is None:
            account_id = _account_id_from_row(row)

    if header_index is None:
        if emit_no_header:
            issues.append(ImportIssue("missing_transaction_header", "No recognizable transaction table header was found.", source))
        return ParsedBankStatement(source=source, transactions=(), bank_account_id=account_id, issues=tuple(issues))

    header_currency = None
    for header in materialized[header_index][0]:
        normalized_header = normalize_label(header)
        if normalized_header.startswith(("amount ", "transaction amount ")):
            header_currency = normalize_currency(normalized_header)
            if header_currency:
                break

    for row, ref in materialized[header_index + 1 :]:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        joined = " ".join(str(value or "") for value in row)
        if any(normalize_label(label) in normalize_label(joined) for labels in _BALANCE_LABELS.values() for label in labels):
            continue
        raw_date = _cell(row, columns.get("transaction_date"))
        raw_amount = _cell(row, columns.get("amount"))
        raw_debit = _cell(row, columns.get("debit"))
        raw_credit = _cell(row, columns.get("credit"))
        raw_running_balance = _cell(row, columns.get("balance"))
        raw_currency = _cell(row, columns.get("currency"))
        parsed_date = _parse_date(raw_date)
        explicit_amount, amount_currency = parse_money(raw_amount)
        debit, debit_currency = parse_money(raw_debit)
        credit, credit_currency = parse_money(raw_credit)
        if debit == 0:
            debit = None
        if credit == 0:
            credit = None
        has_raw_amount = raw_amount is not None and bool(str(raw_amount).strip())
        direction = _parse_direction(_cell(row, columns.get("direction")))
        if not has_raw_amount and debit is None and credit is None:
            # A dated row still gets preserved as incomplete; unrelated footer rows do not.
            if parsed_date is None:
                continue
            amount = None
            issues.append(ImportIssue("missing_amount", "Transaction row has no parseable amount.", ref))
        elif has_raw_amount and explicit_amount is None:
            amount = None
            issues.append(ImportIssue("invalid_amount", "Transaction amount could not be parsed.", ref))
        elif debit is not None and credit is not None:
            amount = None
            issues.append(ImportIssue("conflicting_amount_columns", "Row contains both debit and credit amounts.", ref))
        elif debit is not None:
            amount = -abs(debit)
            direction = "debit"
        elif credit is not None:
            amount = abs(credit)
            direction = "credit"
        else:
            amount = explicit_amount
            if direction == "debit":
                amount = -abs(amount) if amount is not None else None
            elif direction == "credit":
                amount = abs(amount) if amount is not None else None
            elif amount is not None and amount < 0:
                direction = "debit"
            elif amount is not None and amount > 0:
                direction = "credit" if raw_amount and re.search(r"[+]|CR\b", str(raw_amount), re.IGNORECASE) else None
                if direction is None:
                    issues.append(ImportIssue("ambiguous_direction", "Unsigned amount has no explicit debit or credit direction.", ref))
        date_value = parsed_date
        if raw_date is not None and date_value is None:
            issues.append(ImportIssue("invalid_date", "Transaction date could not be parsed.", ref))
        if date_value is None and amount is None:
            continue
        if date_value is not None and amount is not None and raw_running_balance is not None:
            running_balance, _ = parse_money(raw_running_balance)
            if running_balance is not None:
                running_rows.append((date_value, amount, running_balance, ref))
        row_account_id = _clean(_cell(row, columns.get("account_id")))
        transaction_account_id = row_account_id or account_id
        if transaction_account_id:
            transaction_account_ids.add(transaction_account_id)
        currency = normalize_currency(raw_currency) or amount_currency or debit_currency or credit_currency or header_currency
        if currency:
            currencies.add(currency)
        counterparty = _clean(_cell(row, columns.get("counterparty")))
        counterparty = counterparty or _clean(_cell(row, columns.get("merchant")))
        description = _clean(_cell(row, columns.get("description")))
        reference = _clean(_cell(row, columns.get("reference")))
        if not reference:
            reference = _reference_from_text(" ".join(str(value or "") for value in row))
        if not description:
            description = _clean(_cell(row, columns.get("counterparty")))
        transactions.append(
            BankTransaction(
                transaction_date=date_value,
                amount=amount,
                currency=currency,
                source=ref,
                bank_account_id=transaction_account_id,
                counterparty=counterparty,
                description=description,
                reference=reference,
                direction=direction,
            )
        )

    currencies.update(currency for _, currency in balances.values() if currency)
    statement_currency = next(iter(currencies)) if len(currencies) == 1 else None
    opening_balance = balances.get("opening", (None, None))[0]
    closing_balance = balances.get("closing", (None, None))[0]
    if running_rows and len(running_rows) == len(transactions) and opening_balance is None and closing_balance is None:
        ordered_rows = (running_rows, list(reversed(running_rows)))
        chronological = next(
            (
                rows
                for rows in ordered_rows
                if all(next_row[2] - current[2] == next_row[1] for current, next_row in pairwise(rows))
            ),
            None,
        )
        if chronological:
            opening_balance = chronological[0][2] - chronological[0][1]
            closing_balance = chronological[-1][2]
            issues.append(ImportIssue(
                "derived_opening_balance",
                "Opening balance is implied by the earliest running balance and movement; confirm it against a bank statement.",
                chronological[0][3],
            ))
        else:
            issues.append(ImportIssue(
                "inconsistent_running_balance",
                "Running balances do not reconcile to adjacent transaction movements.",
                source,
            ))
    if account_id:
        transaction_account_ids.add(account_id)
    statement_account_id = next(iter(transaction_account_ids)) if len(transaction_account_ids) == 1 else None
    return ParsedBankStatement(
        source=source,
        transactions=tuple(transactions),
        opening_balance=opening_balance,
        closing_balance=closing_balance,
        currency=statement_currency,
        bank_account_id=statement_account_id,
        issues=tuple(issues),
    )


def _parse_pdf_transaction(
    line: str,
    source: SourceReference,
    account_id: str | None,
    issues: list[ImportIssue],
) -> BankTransaction | None:
    date_match = _DATE_AT_START.match(line)
    if not date_match:
        return None
    transaction_date = _parse_date(date_match.group(1))
    remainder = line[date_match.end() :].strip()
    if not remainder or any(normalize_label(label) in normalize_label(line) for labels in _BALANCE_LABELS.values() for label in labels):
        return None
    reference_match = _REFERENCE_LABEL.search(remainder)
    reference = reference_match.group(1) if reference_match else None
    amount_text = remainder[: reference_match.start()].rstrip() if reference_match else remainder
    # For text-only PDFs, accept a terminal amount only with cents, a sign, or currency.
    amount_match = re.search(
        rf"(?P<prefix>.*?)(?P<amount>\(?[+-]?(?:\d[\d ,.']*\d|\d)(?:[.,]\d{{1,2}})?\)?\s*{_CURRENCY_CODE}?)\s*$",
        amount_text,
        re.IGNORECASE,
    )
    amount: Decimal | None = None
    currency: str | None = None
    description = amount_text
    explicit_credit_sign = False
    if amount_match:
        raw_amount = amount_match.group("amount").strip()
        numeric_text = re.sub(r"(?i)(EUR|USD|GBP|SEK|NOK|DKK|CHF|PLN|JPY|CAD|AUD|[€£¥$])", "", raw_amount).strip()
        has_cents = bool(re.search(r"[,.]\d{1,2}\)?$", numeric_text))
        if has_cents or raw_amount.startswith(("+", "-", "(")) or re.search(rf"(?i){_CURRENCY_CODE}", raw_amount):
            amount, currency = parse_money(raw_amount)
            explicit_credit_sign = raw_amount.startswith("+")
            if amount is not None:
                description = amount_match.group("prefix").strip()
    if amount is None:
        issues.append(ImportIssue("unparsed_pdf_transaction", "A dated PDF row was found but its amount was ambiguous.", source))
    if amount is not None and amount > 0 and not re.search(r"\+", amount_text):
        issues.append(ImportIssue("ambiguous_direction", "Unsigned amount has no explicit debit or credit direction.", source))
    return BankTransaction(
        transaction_date=transaction_date,
        amount=amount,
        currency=currency,
        source=source,
        bank_account_id=account_id,
        description=description or None,
        reference=reference,
        direction="debit" if amount is not None and amount < 0 else ("credit" if amount is not None and amount > 0 and explicit_credit_sign else None),
    )


def _capture_text_metadata(
    line: str,
    row_number: int,
    page_ref: SourceReference,
    balances: dict[str, tuple[Decimal, str | None]],
    issues: list[ImportIssue],
) -> None:
    normalized = normalize_label(line)
    for kind, labels in _BALANCE_LABELS.items():
        if any(normalize_label(label) in normalized for label in labels):
            amount, currency = _find_amount_in_text(line)
            if amount is not None:
                _store_balance(balances, kind, amount, currency, page_ref.at(row_number=row_number), issues)


def _merge_balances(
    balances: dict[str, tuple[Decimal, str | None]],
    parsed: ParsedBankStatement,
    issues: list[ImportIssue],
    source: SourceReference,
) -> None:
    for kind, amount in (("opening", parsed.opening_balance), ("closing", parsed.closing_balance)):
        if amount is not None:
            _store_balance(balances, kind, amount, parsed.currency, source, issues)


def _store_balance(
    balances: dict[str, tuple[Decimal, str | None]],
    kind: str,
    amount: Decimal,
    currency: str | None,
    source: SourceReference,
    issues: list[ImportIssue],
) -> None:
    previous = balances.get(kind)
    if previous and previous[0] != amount:
        issues.append(ImportIssue(f"conflicting_{kind}_balance", f"Conflicting {kind} balances were found.", source))
        return
    balances[kind] = (amount, currency or (previous[1] if previous else None))


def _find_amount_in_row(row: Sequence[object]) -> tuple[Decimal | None, str | None]:
    joined = " ".join(str(value or "") for value in row)
    normalized = normalize_label(joined)
    for labels in _BALANCE_LABELS.values():
        if any(normalize_label(label) in normalized for label in labels):
            for index, cell in enumerate(row):
                if not any(normalize_label(label) in normalize_label(cell) for label in labels):
                    continue
                for candidate in row[index + 1 :]:
                    amount, currency = parse_money(candidate)
                    if amount is not None:
                        return amount, currency
            return _find_amount_in_text(joined)
    for cell in row:
        amount, currency = parse_money(cell)
        if amount is not None:
            return amount, currency
    return None, None


def _find_amount_in_text(text: str) -> tuple[Decimal | None, str | None]:
    matches = re.findall(
        rf"(?i){_CURRENCY_CODE}\s*\(?[+-]?\s*\d[\d ,.']*(?:[.,]\d{{1,2}})?\)?|\(?[+-]?\s*\d[\d ,.']*(?:[.,]\d{{1,2}})?\)?\s*{_CURRENCY_CODE}",
        text,
    )
    for match in reversed(matches):
        amount, currency = parse_money(match)
        if amount is not None:
            return amount, currency
    return None, None


def _find_columns(row: Sequence[object]) -> dict[str, int]:
    normalized = [normalize_label(value) for value in row]
    result: dict[str, int] = {}
    for index, label in enumerate(normalized):
        for field, aliases in _HEADERS.items():
            if label in aliases or (field == "amount" and label.startswith(("amount ", "transaction amount "))):
                result.setdefault(field, index)
    return result


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    for date_format in _DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def _parse_direction(value: object) -> str | None:
    normalized = normalize_label(value)
    if not normalized:
        return None
    for alias, direction in _DIRECTION_ALIASES.items():
        if normalized == alias or normalized.startswith(alias + " "):
            return direction
    return None


def _reference_from_text(text: str) -> str | None:
    match = _REFERENCE_LABEL.search(text)
    return match.group(1) if match else None


def _account_id_from_row(row: Sequence[object]) -> str | None:
    for index, value in enumerate(row[:-1]):
        label = normalize_label(value)
        if label in {"iban", "account number", "bank account", "kontonumber", "konto number"}:
            candidate = str(row[index + 1] or "").strip()
            if candidate:
                return re.sub(r"\s+", "", candidate).upper()
    joined = " ".join(str(value or "") for value in row)
    match = _ACCOUNT_LABEL.search(joined)
    return re.sub(r"\s+", "", match.group(1)).upper() if match else None


def _cell(row: Sequence[object], index: int | None) -> object | None:
    return row[index] if index is not None and index < len(row) else None


def _clean(value: object | None) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None
