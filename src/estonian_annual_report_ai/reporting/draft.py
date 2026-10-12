"""Build a review-only annual-report draft from deterministic statements."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

from estonian_annual_report_ai.accounting import (
    Ledger,
    balance_sheet,
    income_statement,
    validate_annual_report,
)

DEFAULT_MISSING_INFORMATION = (
    "Confirm reporting category and applicable Estonian reporting framework.",
    "Confirm accounting policies, accruals, depreciation, and tax balances.",
    "Confirm related-party transactions, guarantees, and contingent liabilities.",
    "Confirm employee count, management loans/advances, and post-reporting events.",
    "Review source documents and opening balances against the prior annual report.",
)


@dataclass(frozen=True, slots=True)
class DraftIssue:
    code: str
    message: str
    severity: str


@dataclass(frozen=True, slots=True)
class AnnualReportDraft:
    company_name: str
    registry_code: str
    period_start: date
    period_end: date
    status: str
    figure_status: str
    balance_sheet: dict[str, Any] | None
    income_statement: dict[str, Any] | None
    notes_draft: tuple[str, ...]
    reconciliation_issues: tuple[DraftIssue, ...]
    missing_information: tuple[str, ...]
    disclaimer: str

    def to_dict(self) -> dict[str, Any]:
        result = asdict(self)
        result["period_start"] = self.period_start.isoformat()
        result["period_end"] = self.period_end.isoformat()
        return result


def _lines(items: tuple[Any, ...]) -> list[dict[str, str]]:
    return [
        {"account_code": item.account_code, "name": item.name, "amount": str(item.amount)}
        for item in items
    ]


def _excel_text(value: str) -> str:
    """Keep untrusted text beginning with spreadsheet formula markers literal."""
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def build_annual_report_draft(
    *,
    company_name: str,
    registry_code: str,
    period_start: date,
    period_end: date,
    ledger: Ledger,
    reconciliation_issues: tuple[DraftIssue, ...] = (),
    missing_information: tuple[str, ...] = DEFAULT_MISSING_INFORMATION,
    confirmed_figure_references: tuple[str, ...] = (),
) -> AnnualReportDraft:
    """Builds statements only when critical accounting checks are clear.

    Ledger-derived figures remain labelled as calculated; source approval is
    not inferred from arithmetic or from the presence of a source reference.
    """
    validation = validate_annual_report(ledger, period_start, period_end)
    issues = list(reconciliation_issues)
    issues.extend(
        DraftIssue(item.code, item.message, item.severity)
        for item in validation.issues
    )
    blocked = any(issue.severity in ("critical", "error") for issue in issues)
    if blocked:
        balance = None
        income = None
        status = "blocked"
    else:
        bs = balance_sheet(ledger, period_end, period_start)
        pl = income_statement(ledger, period_start, period_end)
        balance = {
            "assets": _lines(bs.assets),
            "liabilities": _lines(bs.liabilities),
            "equity": _lines(bs.equity),
            "total_assets": str(bs.total_assets),
            "total_liabilities": str(bs.total_liabilities),
            "total_equity": str(bs.total_equity),
        }
        income = {
            "revenue": _lines(pl.revenue),
            "expenses": _lines(pl.expenses),
            "total_revenue": str(pl.total_revenue),
            "total_expenses": str(pl.total_expenses),
            "net_profit": str(pl.net_profit),
        }
        status = "ready_for_user_review"

    figure_status = (
        "calculated_from_ledger; some figures have user-confirmed source references"
        if confirmed_figure_references
        else "calculated_from_ledger; not source-verified"
    )
    notes = (
        "Accounting policies and framework: require user confirmation.",
        "Related parties, guarantees, contingencies, and management transactions: require confirmation.",
        "Subsequent events and employee information: require confirmation.",
    )
    return AnnualReportDraft(
        company_name=company_name,
        registry_code=registry_code,
        period_start=period_start,
        period_end=period_end,
        status=status,
        figure_status=figure_status,
        balance_sheet=balance,
        income_statement=income,
        notes_draft=notes,
        reconciliation_issues=tuple(issues),
        missing_information=tuple(missing_information),
        disclaimer=(
            "Draft for review only. Not an Estonian statutory report, tax opinion, "
            "or filing-ready document. No estimates are included in financial statements."
        ),
    )


def export_annual_report_draft(draft: AnnualReportDraft, output_dir: Path) -> dict[str, str]:
    """Write local JSON, XLSX, and PDF draft files; never log their contents."""
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "annual-report-draft.json"
    json_path.write_text(json.dumps(draft.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")

    from openpyxl import Workbook
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Company", _excel_text(draft.company_name)])
    summary.append(["Registry code", _excel_text(draft.registry_code)])
    summary.append(["Period", _excel_text(f"{draft.period_start} to {draft.period_end}")])
    summary.append(["Draft status", _excel_text(draft.status)])
    summary.append(["Figure status", _excel_text(draft.figure_status)])
    summary.append(["Disclaimer", _excel_text(draft.disclaimer)])
    for title, statement in (("Balance Sheet", draft.balance_sheet), ("Income Statement", draft.income_statement)):
        sheet = workbook.create_sheet(title)
        if statement is None:
            sheet.append(["Blocked: unresolved critical reconciliation issue"])
            continue
        for key, value in statement.items():
            sheet.append([key])
            if isinstance(value, list):
                sheet.append(["Account code", "Name", "Amount (EUR)"])
                for item in value:
                    sheet.append([
                        _excel_text(item["account_code"]),
                        _excel_text(item["name"]),
                        Decimal(item["amount"]),
                    ])
            else:
                sheet.append([Decimal(value)])
    missing_sheet = workbook.create_sheet("Missing information")
    for item in draft.missing_information:
        missing_sheet.append([_excel_text(item)])
    xlsx_path = output_dir / "annual-report-draft.xlsx"
    workbook.save(xlsx_path)

    pdf_path = output_dir / "annual-report-draft.pdf"
    page = canvas.Canvas(str(pdf_path), pagesize=A4)
    width, height = A4
    y = height - 48
    page.setTitle("Annual Report Draft")
    page.setFont("Helvetica-Bold", 14)
    page.drawString(48, y, "Annual Report Draft — For Review")
    y -= 24
    page.setFont("Helvetica", 10)
    text_lines = [
        f"Company: {draft.company_name} ({draft.registry_code})",
        f"Period: {draft.period_start} to {draft.period_end}",
        f"Status: {draft.status}",
        f"Figures: {draft.figure_status}",
    ]
    if draft.balance_sheet is None:
        text_lines.extend(["", "Statements withheld because critical issues remain unresolved."])
    else:
        text_lines.extend(["", "Balance Sheet (EUR)"])
        text_lines.extend(f"{k}: {v}" for k, v in draft.balance_sheet.items() if not isinstance(v, list))
        text_lines.extend(["", "Income Statement (EUR)"])
        text_lines.extend(f"{k}: {v}" for k, v in draft.income_statement.items() if not isinstance(v, list))
    text_lines.extend(["", "Missing information:"])
    text_lines.extend(f"- {item}" for item in draft.missing_information)
    text_lines.extend(["", draft.disclaimer])
    # Wrap instead of truncating, so no missing-information item is cut off.
    wrapped = [part for line in text_lines for part in (simpleSplit(line, "Helvetica", 10, width - 96) or [""])]
    for line in wrapped:
        if y < 48:
            page.showPage()
            y = height - 48
            page.setFont("Helvetica", 10)
        page.drawString(48, y, line)
        y -= 15
    page.save()

    return {"json": str(json_path), "xlsx": str(xlsx_path), "pdf": str(pdf_path)}
